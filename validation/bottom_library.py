"""Bottom reflectance library from in situ measurements (model/data/bottom_measured.csv).

Sources: underwater reflectance spectra measured in situ with an Ocean Optics USB2000 (probe
about 5 cm from the target, dive torch, Spectralon panel), all on PANGAEA under CC BY 3.0:

  Heron Reef      Roelfsema & Phinn (2012), doi:10.1594/PANGAEA.804589
  Moreton Bay     Roelfsema, Phinn & Joyce (2016), doi:10.1594/PANGAEA.864310 (underwater only)
  Fiji            Roelfsema & Phinn (2017), doi:10.1594/PANGAEA.872506
  Cook Islands    Roelfsema & Phinn (2017), doi:10.1594/PANGAEA.872505
  Glovers Reef    Roelfsema & Phinn (2013), doi:10.1594/PANGAEA.824861

Download them into $PHOTIC_VALIDATION_DATA/raw/bottom/measured (`python bottom_library.py fetch`).

Each spectrum is averaged into 1 nm bins over 400-800 nm and packed, with its target name and
source, into model/data/bottom_spectra_insitu.nc. For the library it is then smoothed (11 nm
boxcar), sampled every 5 nm over 410-700 nm (the range the
spectrometer and the water path between probe and target leave usable), and put in one class
by its target name (CLASSES; mixed, dead-and-bleached, animal and unidentified targets are left
out). Each class is split in two, alternating in the order of a hash of the spectra's identities:
  - the library half gives the class's spectrum: the median of the shapes (each spectrum
    divided by its reflectance at 550 nm) times the median reflectance at 550 nm;
  - the brightness bounds are the 5th and 95th percentiles of the reflectance at 550 nm over
    the whole class (a class has as few as 20 spectra, too few to split for the bounds);
  - the held-out half (half = 1 in the packed file) is independent truth for synthetic scenes
    (tier 3's _insitu scenarios), so the benchmark does not test the library on its own spectra.

  python bottom_library.py fetch      download the five datasets
  python bottom_library.py pack       pack every spectrum, with its class and half, into
                                      model/data/bottom_spectra_insitu.nc (in git)
  python bottom_library.py build      from the packed file: write the library, print a summary
"""

import collections
import hashlib
import os
import re
import sys
import urllib.request

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DATA = os.environ.get("PHOTIC_VALIDATION_DATA", os.path.expanduser("~/photic_dev/validation_data"))
MEASURED = os.path.join(DATA, "raw", "bottom", "measured")
OUT_LIBRARY = os.path.join(REPO, "model", "data", "bottom_measured.csv")

DATASETS = {  # PANGAEA id: (file, archive member or None)
    804589: ("pangaea_804589.tab", None),
    864310: ("pangaea_864310.xlsx", None),
    872506: ("pangaea_872506.zip", "Fiji_Spectra/Fiji_Spectra.xlsx"),
    872505: ("pangaea_872505.zip", "Cook_Island_Spectra/Cook_Island_Spectra.xlsx"),
    824861: ("pangaea_824861.tab", None),
}

LAM = np.arange(410.0, 700.1, 5.0)
SMOOTH = 5.0  # half-width of the boxcar (nm)
NORM = 550.0

# Classes in the library's order: photic's first three types stay sand, seagrass and coral.
# Each target name is lower-cased and tested against the rules in order; EXCLUDE first.
EXCLUDE = r"dead|bleach|sponge|clam|shell|eel|holothur|zoanthid|lykia|sea rod|gorgon|unident|unknown|" \
          r"cyano|lyngbya|cca|coral+ine|calcareous|soft coral|damsel|dark pixel|#ref|other|white panel|" \
          r"\+| and |with sand|sediment\b.*algae|algae.*sediment|lobophora sediment|caulerpa sediment|silt on|" \
          r"encrusting algae|ca branching|benthic microalgae"
CLASSES = [
    ("sand", r"^sand( light| course| light \(white\)| light white)?$|white sand"),
    ("seagrass", r"halodu|halophila|syringodium|thalassia|zostera|cymodocea|cymadosea|seagrass"),
    ("coral", r"coral|porites|acropora|montip|pocillop|favi|favites|montastrea|agaricia|diplora|millepora|"
              r"brain|mushroom|fire|meandrin|massive|digitate|branching|table|encrust"),
    ("dark_sediment", r"sand dark|dark sand|brown sand|mpb|microphyto|silt|^sediment|mud|sand bma|sand pmb"),
    ("brown_algae", r"sargassum|lobophora|dictyota|dictioyta|padina|turbinaria|hydroclath|brown alga"),
    ("green_algae", r"halimeda|caulerpa|cauelerpa|chlorodesm|boodlea|enteromorpha|acetabularia|dictyosphaeria|"
                    r"udotea|green alga"),
    ("turf_rubble", r"turf|rubble|^rock|pavement"),
]


def fetch():
    os.makedirs(MEASURED, exist_ok=True)
    for pid, (fname, _) in DATASETS.items():
        path = os.path.join(MEASURED, fname)
        if os.path.exists(path):
            continue
        print("PANGAEA.%d" % pid, flush=True)
        urllib.request.urlretrieve("https://doi.pangaea.de/10.1594/PANGAEA.%d?format=textfile" % pid, path + ".part")
        os.replace(path + ".part", path)


def classify(name):
    low = " ".join(str(name).lower().replace("_", " ").split())
    if re.search(EXCLUDE, low):
        return None
    for cls, rule in CLASSES:
        if re.search(rule, low):
            return cls
    return None


def resample(wl, R):
    """Boxcar-smoothed reflectance on LAM; NaN where the spectrum has too few valid samples."""
    wl, R = np.asarray(wl, float), np.asarray(R, float)
    ok = np.isfinite(R) & (wl > 404.0)  # the spectrometers read 0 at the first nanometres
    out = np.full(LAM.shape, np.nan)
    for i, lam in enumerate(LAM):
        m = ok & (np.abs(wl - lam) <= SMOOTH)
        if m.sum() >= 3:
            out[i] = R[m].mean()
    return out


def tab_spectra(path, key_cols, src):
    """PANGAEA tab files: one row per wavelength, spectra keyed by key_cols, target in Comment."""
    txt = open(path, encoding="utf-8").read()
    rows = txt[txt.index("*/") + 2:].lstrip("\n").splitlines()
    head = rows[0].split("\t")
    iw, ir, ic = head.index("Lambda [nm]"), head.index("Refl tot"), head.index("Comment")
    spectra = collections.OrderedDict()
    for r in rows[1:]:
        f = r.split("\t")
        e = spectra.setdefault(tuple(f[k] for k in key_cols), dict(target=f[ic].split(" - ")[0], wl=[], R=[]))
        e["wl"].append(float(f[iw]))
        e["R"].append(float(f[ir]) if f[ir] else np.nan)
    return [dict(src=src, id="/".join(k), target=e["target"], wl=e["wl"], R=e["R"]) for k, e in spectra.items()]


def xlsx_rows(path, member):
    import openpyxl
    if member:
        import io
        import zipfile
        return list(openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(path).read(member)), read_only=True,
                                           data_only=True).worksheets[0].iter_rows(values_only=True))
    return list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))


def uq_xlsx_spectra(path, member, src):
    """Fiji and Cook Islands workbooks: one column per target, wavelengths in column 0."""
    rows = xlsx_rows(path, member)
    start = next(i for i, r in enumerate(rows) if r[0] == "Wavelength (below)") + 1
    data = [r for r in rows[start:] if isinstance(r[0], (int, float))]
    wl = [r[0] for r in data]
    out = []
    for c in range(1, len(rows[0])):
        if rows[0][c] is None:
            continue
        R = [r[c] if c < len(r) and isinstance(r[c], (int, float)) else np.nan for r in data]
        out.append(dict(src=src, id="%s/%d" % (src, c), target=str(rows[0][c]), wl=wl, R=R))
    return out


def moreton_spectra(path):
    rows = xlsx_rows(path, None)
    wl = [r[0] for r in rows[17:]]
    out = []
    for c in range(2, len(rows[16])):
        name, place, fov = rows[16][c], str(rows[0][c]), str(rows[5][c])
        if name is None or "uw" not in fov or "Roof" in place:
            continue  # measured in air, or on the roof
        R = [r[c] if isinstance(r[c], (int, float)) else np.nan for r in rows[17:]]
        out.append(dict(src="Moreton Bay", id="Moreton/%d" % c, target=str(name), wl=wl, R=R))
    return out


SOURCES = [  # name, PANGAEA id, citation
    ("Heron Reef", 804589, "Roelfsema, C. M. and Phinn, S. R. (2012): Spectral reflectance library of selected biotic "
     "and abiotic coral reef features in Heron Reef. PANGAEA, https://doi.org/10.1594/PANGAEA.804589"),
    ("Moreton Bay", 864310, "Roelfsema, C. M., Phinn, S. R. and Joyce, K. (2016): Spectral reflectance library of algal, "
     "seagrass and substrate types in Moreton Bay, Australia. PANGAEA, https://doi.org/10.1594/PANGAEA.864310"),
    ("Fiji", 872506, "Roelfsema, C. M. and Phinn, S. R. (2017): Spectral reflectance library of healthy corals, "
     "bleached corals and other benthic features in Fiji. PANGAEA, https://doi.org/10.1594/PANGAEA.872506"),
    ("Cook Islands", 872505, "Roelfsema, C. M. and Phinn, S. R. (2017): Spectral reflectance library of corals and "
     "benthic features in Aitutaki, Cook Islands. PANGAEA, https://doi.org/10.1594/PANGAEA.872505"),
    ("Glovers Reef", 824861, "Roelfsema, C. M. and Phinn, S. R. (2013): Spectral reflectance library of selected biotic "
     "and abiotic coral reef features in Glovers Reef, Belize. PANGAEA, https://doi.org/10.1594/PANGAEA.824861"),
]
PACKED = os.path.join(REPO, "model", "data", "bottom_spectra_insitu.nc")
GRID = np.arange(400.0, 800.1, 1.0)  # the packed file's wavelengths


def read_sources():
    """Every spectrum in the five downloads, on its native wavelengths."""
    s = []
    s += tab_spectra(os.path.join(MEASURED, DATASETS[804589][0]), (0,), "Heron Reef")
    s += moreton_spectra(os.path.join(MEASURED, DATASETS[864310][0]))
    s += uq_xlsx_spectra(os.path.join(MEASURED, DATASETS[872506][0]), DATASETS[872506][1], "Fiji")
    s += uq_xlsx_spectra(os.path.join(MEASURED, DATASETS[872505][0]), DATASETS[872505][1], "Cook Islands")
    s += tab_spectra(os.path.join(MEASURED, DATASETS[824861][0]), (0, 5), "Glovers Reef")
    return [e for e in s if e["target"].strip() and e["target"] not in ("#REF!", "Dark Pixel")]


def bin_1nm(wl, R):
    """Mean of the native samples in each 1 nm bin of GRID (NaN where none, or where the
    spectrometer reads 0 below 405 nm)."""
    wl, R = np.asarray(wl, float), np.asarray(R, float)
    ok = np.isfinite(R) & ~((wl < 405.0) & (R == 0.0))
    k = np.rint(wl[ok] - GRID[0]).astype(int)
    m = (k >= 0) & (k < len(GRID))
    total = np.bincount(k[m], weights=R[ok][m], minlength=len(GRID))
    count = np.bincount(k[m], minlength=len(GRID))
    with np.errstate(invalid="ignore"):
        return np.where(count > 0, total/np.maximum(count, 1), np.nan)


def pack():
    """Write every in situ spectrum, on 1 nm bins over 400-800 nm, with its class and half, to
    model/data/bottom_spectra_insitu.nc."""
    import netCDF4
    spectra = read_sources()
    R = np.array([bin_1nm(e["wl"], e["R"]) for e in spectra])
    cls, half = [], []
    for e, r in zip(spectra, R):
        c = classify(e["target"])
        usable = None
        if c is not None:
            Rs = resample(GRID, r)
            r550 = np.interp(NORM, LAM, Rs)
            # Usable: complete over 410-700 nm, physical, and not so dark that noise dominates.
            usable = np.all(np.isfinite(Rs)) and 0.01 <= r550 <= 1.0 and Rs.min() >= -0.005
        cls.append(c if usable else "")
        half.append(-1)
    # Within each class, alternate in the order of a hash of the identities: 0 library, 1 held out.
    for c, _ in CLASSES:
        idx = sorted((i for i in range(len(spectra)) if cls[i] == c),
                     key=lambda i: hashlib.sha1(spectra[i]["id"].encode()).hexdigest())
        for n, i in enumerate(idx):
            half[i] = n % 2
    src_index = {name: k for k, (name, _, _) in enumerate(SOURCES)}
    with netCDF4.Dataset(PACKED, "w") as nc:
        nc.title = "In situ underwater reflectance of benthic types from five coral reef and seagrass libraries"
        nc.summary = ("Reflectance relative to a Spectralon panel, measured underwater with an Ocean Optics USB2000 "
                      "spectrometer and an artificial light, probe about 5 cm from the target. Averaged into 1 nm "
                      "bins; NaN where a source has no sample. class and half are photic's (validation/"
                      "bottom_library.py): the class from the target name, and the library (0) or held-out (1) "
                      "half; empty class and half -1 for targets left out.")
        nc.license = "CC BY 3.0 (https://creativecommons.org/licenses/by/3.0/); cite the sources below"
        nc.sources = "\n".join(c for _, _, c in SOURCES)
        nc.history = "validation/bottom_library.py pack"
        nc.createDimension("wavelength", len(GRID))
        nc.createDimension("spectrum", len(spectra))
        nc.createDimension("source", len(SOURCES))
        v = nc.createVariable("wavelength", "f4", ("wavelength",))
        v.units = "nm"
        v[:] = GRID
        v = nc.createVariable("reflectance", "f4", ("spectrum", "wavelength"), zlib=True, complevel=9,
                              fill_value=np.float32(np.nan))
        v.units = "1"
        v.long_name = "underwater bottom reflectance factor"
        v[:, :] = R.astype(np.float32)
        for name, values, doc in (("target", [e["target"] for e in spectra], "target name as in the source"),
                                  ("id", [e["id"] for e in spectra], "spectrum identifier within the source"),
                                  ("class", cls, "photic class (empty: not used)")):
            v = nc.createVariable(name, str, ("spectrum",))
            v.long_name = doc
            v[:] = np.array(values, dtype=object)
        v = nc.createVariable("source", "i1", ("spectrum",))
        v.long_name = "index into source_name"
        v[:] = [src_index[e["src"]] for e in spectra]
        v = nc.createVariable("half", "i1", ("spectrum",))
        v.long_name = "0 library, 1 held out, -1 not used"
        v[:] = half
        v = nc.createVariable("source_name", str, ("source",))
        v[:] = np.array([n for n, _, _ in SOURCES], dtype=object)
        v = nc.createVariable("source_doi", str, ("source",))
        v[:] = np.array(["10.1594/PANGAEA.%d" % p for _, p, _ in SOURCES], dtype=object)
    print("%d spectra (%d classified) -> %s, %.2f MB" % (len(spectra), sum(1 for c in cls if c), PACKED,
                                                      os.path.getsize(PACKED)/1e6))


def load_all():
    """The classified spectra from the packed file, smoothed onto LAM."""
    import netCDF4
    with netCDF4.Dataset(PACKED) as nc:
        R = np.array(nc["reflectance"][:].filled(np.nan), float)
        cls, half, ids, targets = nc["class"][:], nc["half"][:], nc["id"][:], nc["target"][:]
        names = nc["source_name"][:]
        src = nc["source"][:]
    out = []
    for i in range(len(R)):
        if not cls[i]:
            continue
        Rs = resample(GRID, R[i])
        out.append(dict(src=str(names[src[i]]), id=str(ids[i]), target=str(targets[i]), cls=str(cls[i]),
                        half=int(half[i]), Rs=np.clip(Rs, 0.0, None), r550=float(np.interp(NORM, LAM, Rs))))
    return out


def build():
    spectra = load_all()
    lib, stats = [], []
    for cls, _ in CLASSES:
        members = [e for e in spectra if e["cls"] == cls]
        part = [e for e in members if e["half"] == 0]
        if len(part) < 3:
            sys.exit("class %s: only %d library spectra" % (cls, len(part)))
        shapes = np.array([e["Rs"]/e["r550"] for e in part])
        r550 = np.array([e["r550"] for e in part])
        R = np.median(shapes, axis=0)*np.median(r550)
        lo, hi = np.percentile([e["r550"] for e in members], [5, 95])
        lib.append((cls, R, lo, hi))
        stats.append((cls, len(members), len(part), collections.Counter(e["src"] for e in part), np.median(r550), lo, hi))

    with open(OUT_LIBRARY, "w") as f:
        f.write("# Bottom reflectance library from in situ measurements: per class, the median shape times\n"
                "# the median reflectance at 550 nm of half the class's spectra (the other half is held out\n"
                "# for synthetic tests); brightness_min/brightness_max are the 5th/95th percentiles of the\n"
                "# reflectance at 550 nm over the whole class. Built by validation/bottom_library.py from\n"
                "# model/data/bottom_spectra_insitu.nc.\n"
                "# Reflectance relative to a Spectralon panel, measured underwater with an Ocean Optics\n"
                "# USB2000 and a dive torch, probe ~5 cm from the target;\n"
                "# smoothed with an 11 nm boxcar, 410-700 nm every 5 nm. Sources (CC BY 3.0, PANGAEA):\n"
                "#   Roelfsema & Phinn (2012), Heron Reef, doi:10.1594/PANGAEA.804589\n"
                "#   Roelfsema, Phinn & Joyce (2016), Moreton Bay, doi:10.1594/PANGAEA.864310\n"
                "#   Roelfsema & Phinn (2017), Fiji, doi:10.1594/PANGAEA.872506\n"
                "#   Roelfsema & Phinn (2017), Cook Islands (Aitutaki), doi:10.1594/PANGAEA.872505\n"
                "#   Roelfsema & Phinn (2013), Glovers Reef, Belize, doi:10.1594/PANGAEA.824861\n"
                "# Spectra per class (library half, by source):\n")
        for cls, n, nl, src, med, lo, hi in stats:
            f.write("#   %-14s %3d of %3d: %s\n" % (cls, nl, n, ", ".join("%s %d" % kv for kv in sorted(src.items()))))
        f.write("# Bands outside 410-700 nm are extrapolated linearly by photic; do not use them.\n")
        f.write("wavelength," + ",".join(c for c, *_ in lib) + "\n")
        for i, lam in enumerate(LAM):
            f.write("%g," % lam + ",".join("%.5f" % R[i] for _, R, _, _ in lib) + "\n")
        f.write("brightness_min," + ",".join("%.3f" % lo for _, _, lo, _ in lib) + "\n")
        f.write("brightness_max," + ",".join("%.3f" % hi for _, _, _, hi in lib) + "\n")


    print("%-14s %5s %5s %7s %13s | shape at 443/490/560/665 (/550)" % ("class", "n", "lib", "R550", "bounds"))
    for (cls, R, lo, hi), (_, n, nl, src, med, _, _) in zip(lib, stats):
        sh = np.interp([443, 490, 560, 665], LAM, R)/np.interp(NORM, LAM, R)
        print("%-14s %5d %5d %7.3f %6.3f-%.3f | %s" % (cls, n, nl, med, lo, hi, " ".join("%.2f" % x for x in sh)))
    print("held out: %d spectra (tier 3's _insitu bottoms)" % sum(e["half"] == 1 for e in spectra))
    print("library -> %s" % OUT_LIBRARY)


if __name__ == "__main__":
    {"fetch": fetch, "pack": pack, "build": build}[sys.argv[1]]()
